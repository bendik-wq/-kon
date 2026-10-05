"""Gjør eksempeloppgavene i Excel om til JSON som nettsiden leser.

Bruk (fra repo-roten):
    python verktoy/xlsx_til_json.py

Leser   eksempler/HSM122_eksempeloppgaver.xlsx
Skriver eksempler/oppgaver.json

Cellenes rolle leses fra fargekodene i arket:
  gul bakgrunn (FFF2CC)   -> inndata (redigerbar på nettsiden)
  grønn bakgrunn (E2F0D9) -> hovedsvar
  svart rad (1D1D1F)      -> ny seksjon
  lyseblå rad (EEF4FF)    -> kolonneoverskrifter
  grå tekst (6E6E73)      -> forklaring / formel
Lagre arket i Excel før du kjører skriptet, slik at de utregnede verdiene er med.
"""
import json
import os
import sys

from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, 'eksempler', 'HSM122_eksempeloppgaver.xlsx')
OUT = os.path.join(ROOT, 'eksempler', 'oppgaver.json')


def rgb(color):
    if color is None or color.type != 'rgb' or not isinstance(color.rgb, str):
        return None
    return color.rgb[-6:].upper()


def fill_of(cell):
    f = cell.fill
    if f is None or f.fill_type != 'solid':
        return None
    return rgb(f.fgColor)


def font_rgb(cell):
    return rgb(cell.font.color) if cell.font and cell.font.color else None


def clean(v):
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def main():
    wf = load_workbook(SRC)                   # formulas
    wv = load_workbook(SRC, data_only=True)   # cached values (needs a save from Excel)
    sheets = []
    for ws in wf.worksheets:
        if ws.title == 'Oversikt':
            continue
        vs = wv[ws.title]
        max_r, max_c = ws.max_row, ws.max_column
        grid = [[None] * max_c for _ in range(max_r)]
        cache = [[None] * max_c for _ in range(max_r)]
        merged_task = None
        for rng in ws.merged_cells.ranges:
            if fill_of(ws.cell(rng.min_row, rng.min_col)) == 'F5F5F7':
                merged_task = rng.min_row

        for row in ws.iter_rows(min_row=1, max_row=max_r, max_col=max_c):
            for cell in row:
                r, c = cell.row - 1, cell.column - 1
                cache[r][c] = clean(vs.cell(cell.row, cell.column).value)
                if cell.data_type == 'f':
                    grid[r][c] = str(cell.value)
                elif isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                    grid[r][c] = clean(cell.value)

        sheet = {
            'name': ws.title,
            'title': ws.cell(1, 1).value,
            'task': ws.cell(merged_task, 1).value if merged_task else '',
            'rows': max_r, 'cols': max_c,
            'grid': grid, 'cache': cache, 'blocks': [],
        }
        block, table = None, None
        for r in range(3, max_r + 1):
            if r == merged_task:
                continue
            a = ws.cell(r, 1)
            fa = fill_of(a)
            if fa == '1D1D1F':
                block = {'title': a.value, 'parts': []}
                sheet['blocks'].append(block)
                table = None
                continue
            if block is None:
                continue
            if fa == 'EEF4FF':
                heads = [ws.cell(r, c).value for c in range(1, max_c + 1)]
                while heads and heads[-1] in (None, ''):
                    heads.pop()
                table = {'type': 'table', 'cols': [h or '' for h in heads], 'rows': []}
                block['parts'].append(table)
                continue
            cells = []
            for c in range(1, max_c + 1):
                x = ws.cell(r, c)
                if x.value in (None, '') and fill_of(x) != 'FFF2CC':
                    cells.append(None)
                    continue
                role = 'text'
                if fill_of(x) == 'FFF2CC':
                    role = 'input'
                elif x.data_type == 'f':
                    role = 'formula'
                elif font_rgb(x) == '6E6E73':
                    role = 'desc'
                cells.append({'r': r - 1, 'c': c - 1, 'role': role,
                              'fmt': x.number_format if x.number_format != 'General' else None,
                              'bold': bool(x.font and x.font.b),
                              'text': x.value if role in ('text', 'desc') else None})
            while cells and cells[-1] is None:
                cells.pop()
            if not cells:
                continue
            only_a = len(cells) == 1 and cells[0]
            if only_a and cells[0]['role'] == 'desc':
                block['parts'].append({'type': 'note', 'text': cells[0]['text']})
                table = None
                continue
            if only_a and font_rgb(a) == '0071E3':
                if table:
                    table['rows'].append({'kind': 'sub', 'text': a.value})
                continue
            if table is None:
                table = {'type': 'table', 'cols': [], 'rows': []}
                block['parts'].append(table)
            table['rows'].append({'kind': 'row', 'result': fill_of(a) == 'E2F0D9', 'cells': cells})
        sheets.append(sheet)

    with open(OUT, 'w', encoding='utf-8') as fh:
        json.dump({'source': os.path.basename(SRC), 'sheets': sheets}, fh, ensure_ascii=False, separators=(',', ':'))
    print(f'Skrev {OUT} ({len(sheets)} ark)')


if __name__ == '__main__':
    sys.exit(main())
